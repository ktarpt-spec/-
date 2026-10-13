"""収支計画（月次36か月）の試算・グラフ（SVG）・Excelシート生成

usage: finance.py <outdir> [<企画書.xlsx>]
  outdir/drawings/fin_trend.svg, fin_years.svg, fin_cac.svg と build/finance.json を出力。
  xlsx を渡すと「収支計画」シートを追加（数式入り）。
前提（すべて概算・単位 万円/月）:
  ・最終月商 5,600万円（36か月目）。立上げは指数カーブで逓増。
  ・自費診療・自費施術100%のため売上原価（消耗品・MRI保守・読影等）6%、粗利率94%。
  ・賃料・共益費 約900万円/月（最終月商比 約16%）。
"""
import json, math, os, sys

FINAL = 5600
MONTHS = 36
RENT = 900                 # 賃料・共益費（万円/月）
COGS = 0.06                # 売上原価率
AD = {1: 300, 2: 250, 3: 220}           # 広告宣伝費（万円/月）
LABOR = lambda s: 1152 + 0.08 * s       # 人件費（固定＋歩合）
OTHER = lambda s: 600 + 0.02 * s        # その他販管費（水光熱・システム・減価償却等）
INDUSTRY = dict(rent=9, ad=17)          # 一般的な自費クリニック（路面・空中階の平均的立地）の目安（％）


def sales(m):
    if m >= MONTHS:
        return FINAL
    return round(FINAL * (1 - 0.75 * math.exp(-m / 9.0)) / 10) * 10


def model():
    rows = []
    for m in range(1, MONTHS + 1):
        s = sales(m); y = (m - 1) // 12 + 1
        c = dict(month=m, year=y, sales=s, cogs=round(s * COGS), labor=round(LABOR(s)), rent=RENT, ad=AD[y], other=round(OTHER(s)))
        c["sga"] = c["labor"] + c["rent"] + c["ad"] + c["other"]
        c["op"] = s - c["cogs"] - c["sga"]
        rows.append(c)
    years = []
    for y in (1, 2, 3):
        rs = [r for r in rows if r["year"] == y]
        t = {k: sum(r[k] for r in rs) for k in ("sales", "cogs", "labor", "rent", "ad", "other", "sga", "op")}
        t["year"] = y
        years.append(t)
    fin = rows[-1]
    return rows, years, fin


C = dict(s1="#2a78d6", s2="#eb6834", s3="#1baf7a", surf="#fcfcfb", tp="#0b0b0b", ts="#52514e", mut="#898781", grid="#e1e0d9", base="#c3c2b7")
FONT = "'IPAGothic',sans-serif"


def svg_wrap(w, h, body):
    return (f'<svg xmlns="http://www.w3.org/2000/svg" width="{w}mm" height="{h}mm" viewBox="0 0 {w} {h}">'
            f'<rect width="{w}" height="{h}" fill="{C["surf"]}"/>{body}</svg>')


def t(x, y, s, size=2.4, fill=C["ts"], anchor="middle", weight="normal"):
    return f'<text x="{x:.2f}" y="{y:.2f}" font-size="{size}" fill="{fill}" text-anchor="{anchor}" font-weight="{weight}" font-family="{FONT}">{s}</text>'


def trend_svg(rows):
    """月商の推移（36か月）＋営業利益（同一単位・同一軸）"""
    W, H = 180, 62
    L, R, T, B = 16, 30, 7, 9
    pw, ph = W - L - R, H - T - B
    ymin, ymax = -500, 6000
    X = lambda m: L + (m - 1) / (MONTHS - 1) * pw
    Y = lambda v: T + (ymax - v) / (ymax - ymin) * ph
    o = []
    for v in range(0, 6001, 1000):
        o.append(f'<line x1="{L}" x2="{L+pw}" y1="{Y(v):.2f}" y2="{Y(v):.2f}" stroke="{C["grid"] if v else C["base"]}" stroke-width="{0.15 if v else 0.3}"/>')
        o.append(t(L - 1.5, Y(v) + 0.8, f"{v:,}", 2.2, C["mut"], "end"))
    for m in (1, 12, 24, 36):
        o.append(t(X(m), T + ph + 4.0, f"{m}か月", 2.2, C["mut"]))
    for yb in (12.5, 24.5):
        xx = X(yb)
        o.append(f'<line x1="{xx:.2f}" x2="{xx:.2f}" y1="{T}" y2="{T+ph}" stroke="{C["grid"]}" stroke-width="0.2" stroke-dasharray="1,1"/>')
    for y, lab in ((1, "1年目"), (2, "2年目"), (3, "3年目")):
        o.append(t(X(6.5 + (y - 1) * 12), T + 2.5, lab, 2.3, C["ts"]))
    for key, col, name in (("sales", C["s1"], "月商"), ("op", C["s3"], "営業利益")):
        pts = " ".join(f"{X(r['month']):.2f},{Y(r[key]):.2f}" for r in rows)
        o.append(f'<polyline points="{pts}" fill="none" stroke="{col}" stroke-width="0.6" stroke-linejoin="round" stroke-linecap="round"/>')
        last = rows[-1]
        o.append(f'<circle cx="{X(36):.2f}" cy="{Y(last[key]):.2f}" r="0.9" fill="{col}" stroke="{C["surf"]}" stroke-width="0.4"/>')
        o.append(t(X(36) + 2.5, Y(last[key]) + 0.9, f"{name} {last[key]:,}万円", 2.4, C["tp"], "start", "bold"))
    # 賃料ライン（参考）
    o.append(f'<line x1="{L}" x2="{L+pw}" y1="{Y(RENT):.2f}" y2="{Y(RENT):.2f}" stroke="{C["s2"]}" stroke-width="0.4" stroke-dasharray="1.6,1"/>')
    o.append(t(X(36) + 2.5, Y(RENT) + 0.9, f"賃料等 {RENT}万円", 2.3, C["tp"], "start"))
    be = next(r["month"] for r in rows if r["op"] > 0)
    o.append(f'<circle cx="{X(be):.2f}" cy="{Y(0):.2f}" r="1.0" fill="none" stroke="{C["tp"]}" stroke-width="0.3"/>')
    o.append(t(X(be) + 1.5, Y(0) + 4.5, f"{be}か月目に単月黒字化", 2.3, C["tp"], "start"))
    o.append(t(L - 1.5, T - 2.5, "万円／月", 2.2, C["mut"], "end"))
    return svg_wrap(W, H, "".join(o)), be


def years_svg(years):
    """年次：売上・費用（原価＋販管費）・営業利益（月平均、同一軸）"""
    W, H = 180, 56
    L, R, T, B = 16, 4, 8, 9
    pw, ph = W - L - R, H - T - B
    ymin, ymax = -500, 6000
    Y = lambda v: T + (ymax - v) / (ymax - ymin) * ph
    o = []
    for v in range(0, 6001, 1000):
        o.append(f'<line x1="{L}" x2="{L+pw}" y1="{Y(v):.2f}" y2="{Y(v):.2f}" stroke="{C["grid"] if v else C["base"]}" stroke-width="{0.15 if v else 0.3}"/>')
        o.append(t(L - 1.5, Y(v) + 0.8, f"{v:,}", 2.2, C["mut"], "end"))
    gw = pw / 3
    bw = 11
    for i, yv in enumerate(years):
        vals = [("売上", yv["sales"] / 12, C["s1"]), ("費用", (yv["cogs"] + yv["sga"]) / 12, C["s2"]), ("営業利益", yv["op"] / 12, C["s3"])]
        x0 = L + gw * i + (gw - 3 * bw - 2 * 1.2) / 2
        for j, (nm, v, col) in enumerate(vals):
            x = x0 + j * (bw + 1.2)
            y1, y2 = Y(max(v, 0)), Y(min(v, 0))
            r_ = 0.8
            if v >= 0:
                o.append(f'<path d="M{x},{y2} L{x},{y1 + r_} Q{x},{y1} {x + r_},{y1} L{x + bw - r_},{y1} Q{x + bw},{y1} {x + bw},{y1 + r_} L{x + bw},{y2} Z" fill="{col}"/>')
                o.append(t(x + bw / 2, y1 - 1.2, f"{v:,.0f}", 2.3, C["tp"]))
            else:
                o.append(f'<rect x="{x}" y="{y1}" width="{bw}" height="{y2 - y1}" fill="{col}"/>')
                o.append(t(x + bw / 2, y2 + 3.0, f"{v:,.0f}", 2.3, C["tp"]))
        o.append(t(L + gw * i + gw / 2, T + ph + 4.5, f"{yv['year']}年目（年商 {yv['sales']/10000:.2f}億円）", 2.4, C["ts"]))
    # 凡例
    lx = L + 2
    for j, (nm, col) in enumerate((("売上（月平均）", C["s1"]), ("費用：原価＋販管費", C["s2"]), ("営業利益", C["s3"]))):
        o.append(f'<rect x="{lx + j*40}" y="{T - 5.5}" width="3" height="2.4" rx="0.5" fill="{col}"/>')
        o.append(t(lx + j * 40 + 4.2, T - 3.6, nm, 2.3, C["ts"], "start"))
    o.append(t(W - R, T - 3.6, "万円／月", 2.2, C["mut"], "end"))
    return svg_wrap(W, H, "".join(o))


def cac_svg(fin):
    """賃料＋広告宣伝費（集客コスト）の比較：一般的な自費クリニック vs 東京駅院（同一軸・水平積み上げ）"""
    W, H = 120, 36
    L, R = 40, 14
    pw = W - L - R
    X = lambda p: L + p / 30 * pw
    rent_p = fin["rent"] / fin["sales"] * 100
    ad_p = fin["ad"] / fin["sales"] * 100
    rows = [("一般的な自費クリニック", INDUSTRY["rent"], INDUSTRY["ad"]), ("東京駅院（最終月）", rent_p, ad_p)]
    o = []
    for i, (nm, r_, a_) in enumerate(rows):
        y = 8 + i * 13
        o.append(t(L - 2, y + 4, nm, 2.5, C["tp"], "end", "bold" if i else "normal"))
        o.append(f'<rect x="{L}" y="{y}" width="{X(r_) - L:.2f}" height="6" fill="{C["s1"]}"/>')
        o.append(f'<rect x="{X(r_) + 0.4:.2f}" y="{y}" width="{X(r_ + a_) - X(r_) - 0.4:.2f}" height="6" fill="{C["s2"]}"/>')
        o.append(t((L + X(r_)) / 2, y + 4.1, f"賃料 {r_:.0f}%", 2.2, "#ffffff"))
        o.append(t((X(r_) + X(r_ + a_)) / 2, y + 4.1, f"広告 {a_:.0f}%", 2.2, C["tp"] if a_ < 6 else "#ffffff"))
        o.append(t(X(r_ + a_) + 1.5, y + 4.1, f"計 {r_ + a_:.0f}%", 2.5, C["tp"], "start", "bold"))
    o.append(f'<line x1="{L}" x2="{L}" y1="5" y2="{H - 4}" stroke="{C["base"]}" stroke-width="0.3"/>')
    o.append(t(L, H - 1, "売上に対する比率（集客コスト＝賃料＋広告宣伝費）", 2.1, C["mut"], "start"))
    return svg_wrap(W, H, "".join(o))


def add_sheet(xlsx, rows):
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    wb = openpyxl.load_workbook(xlsx)
    if "収支計画" in wb.sheetnames:
        del wb["収支計画"]
    ws = wb.create_sheet("収支計画")
    hd = PatternFill("solid", fgColor="F3EDE4"); thin = Side(style="thin", color="C9C0B5")
    ws["A1"] = "収支計画（概算）　整形外科ふっか／整痛院ふっか 東京駅院　単位：万円"; ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "前提：最終月商5,600万円（36か月目）／自費100％・売上原価率6％／賃料・共益費900万円/月（最終月商比 約16％）。人件費=1,152+売上×8％、その他販管費=600+売上×2％。"
    ws["A2"].font = Font(size=9, color="52514E")
    # 年次サマリー（数式）
    ws["A4"] = "年次サマリー"; ws["A4"].font = Font(bold=True)
    head = ["項目", "1年目", "2年目", "3年目"]
    items = [("売上高", "C"), ("売上原価", "D"), ("売上総利益", None), ("人件費", "E"), ("賃料・共益費", "F"), ("広告宣伝費", "G"),
             ("その他販管費", "H"), ("販管費計", None), ("営業利益", None), ("営業利益率", None), ("賃料比率（対売上）", None)]
    for j, h in enumerate(head):
        c = ws.cell(row=5, column=1 + j, value=h); c.fill = hd; c.font = Font(bold=True)
    base = 20  # 月次データ開始行
    for i, (nm, col) in enumerate(items):
        r = 6 + i
        ws.cell(row=r, column=1, value=nm)
        for y in (1, 2, 3):
            a, b = base + (y - 1) * 12, base + y * 12 - 1
            cl = ws.cell(row=r, column=1 + y)
            L_ = openpyxl.utils.get_column_letter(1 + y)
            if col:
                cl.value = f"=SUM({col}{a}:{col}{b})"
            elif nm == "売上総利益":
                cl.value = f"={L_}6-{L_}7"
            elif nm == "販管費計":
                cl.value = f"=SUM({L_}9:{L_}12)"
            elif nm == "営業利益":
                cl.value = f"={L_}8-{L_}13"
            elif nm == "営業利益率":
                cl.value = f"={L_}14/{L_}6"; cl.number_format = "0.0%"
            elif nm == "賃料比率（対売上）":
                cl.value = f"={L_}10/{L_}6"; cl.number_format = "0.0%"
            if cl.number_format == "General":
                cl.number_format = "#,##0"
    # 月次
    ws.cell(row=base - 1, column=1, value="月次").font = Font(bold=True)
    mh = ["年", "月", "売上高", "売上原価", "人件費", "賃料・共益費", "広告宣伝費", "その他販管費", "販管費計", "営業利益", "賃料比率"]
    for j, h in enumerate(mh):
        c = ws.cell(row=base - 1 + 0, column=2 + j - 1 + 1 - 1, value=None)
    for j, h in enumerate(mh):
        c = ws.cell(row=base - 1, column=1 + j, value=h); c.fill = hd; c.font = Font(bold=True)
    ws.cell(row=base - 1, column=1).value = "年"
    for i, rr in enumerate(rows):
        r = base + i
        ws.cell(row=r, column=1, value=f"{rr['year']}年目")
        ws.cell(row=r, column=2, value=rr["month"])
        ws.cell(row=r, column=3, value=rr["sales"])
        ws.cell(row=r, column=4, value=f"=ROUND(C{r}*0.06,0)")
        ws.cell(row=r, column=5, value=f"=ROUND(1152+C{r}*0.08,0)")
        ws.cell(row=r, column=6, value=RENT)
        ws.cell(row=r, column=7, value=rr["ad"])
        ws.cell(row=r, column=8, value=f"=ROUND(600+C{r}*0.02,0)")
        ws.cell(row=r, column=9, value=f"=SUM(E{r}:H{r})")
        ws.cell(row=r, column=10, value=f"=C{r}-D{r}-I{r}")
        ws.cell(row=r, column=11, value=f"=F{r}/C{r}"); ws.cell(row=r, column=11).number_format = "0.0%"
        for c in range(3, 11):
            ws.cell(row=r, column=c).number_format = "#,##0"
    for col, w in zip("ABCDEFGHIJK", (16, 10, 10, 10, 10, 12, 11, 12, 10, 10, 9)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=5, max_row=16, max_col=4):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "19:19"
    wb.save(xlsx)


if __name__ == "__main__":
    out = sys.argv[1]
    rows, years, fin = model()
    tr, be = trend_svg(rows)
    open(os.path.join(out, "drawings", "fin_trend.svg"), "w").write(tr)
    open(os.path.join(out, "drawings", "fin_years.svg"), "w").write(years_svg(years))
    open(os.path.join(out, "drawings", "fin_cac.svg"), "w").write(cac_svg(fin))
    here = os.path.dirname(os.path.abspath(__file__))
    json.dump(dict(rows=rows, years=years, final=fin, breakeven=be), open(os.path.join(here, "build", "finance.json"), "w"), ensure_ascii=False, indent=1)
    if len(sys.argv) > 2:
        add_sheet(sys.argv[2], rows)
    for y in years:
        print(y["year"], "年商", y["sales"], "販管", y["sga"], "営業利益", y["op"], f"利益率 {y['op']/y['sales']*100:.1f}%", f"賃料比率 {y['rent']/y['sales']*100:.1f}%")
    print("final", fin, "breakeven month", be)
