"""貸主様式「借室申込書」(xls→xlsx変換済) に記入する。 usage: fill_app.py <雛形.xlsx> <out.xlsx>
申込人：株式会社ふっか（履歴事項全部証明書より）。1・4・5階のセット賃借。"""
import sys
import openpyxl
from openpyxl.styles import Alignment, Font

FLOORS = [  # 階, ㎡, 坪, 賃料坪単価（円）
    ("1階", 98.36, 29.75, 97000),
    ("4階", 229.86, 69.53, 45000),
    ("5階", 229.86, 69.53, 45000),
]
CAM = 3000            # 共益費（円/坪）
DEPOSIT_MONTHS = 12   # 敷金（月額賃料の月数）


def terms():
    rent = [round(t * u) for _, _, t, u in FLOORS]
    cam = [round(t * CAM) for _, _, t, _ in FLOORS]
    return dict(rent=rent, cam=cam, rent_total=sum(rent), cam_total=sum(cam),
                deposit=sum(rent) * DEPOSIT_MONTHS, tsubo=sum(f[2] for f in FLOORS), m2=sum(f[1] for f in FLOORS))


if __name__ == "__main__":
    src, out = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src)
    ws = wb.active
    T = terms()
    yen = lambda v: f"{v:,}円"

    def put(cell, value, size=None, wrap=True, h="left"):
        c = ws[cell]
        c.value = value
        c.alignment = Alignment(wrap_text=wrap, vertical="center", horizontal=h)
        if size:
            c.font = Font(name=c.font.name, size=size)

    # 1行で収めるセルは結合（雛形は未結合）
    for rg in ("I5:K5", "I7:K7", "I9:J9", "E39:K39", "E40:K40", "E42:K42", "E43:K43", "E44:K44"):
        ws.merge_cells(rg)
    ws.unmerge_cells("E45:K45"); ws.merge_cells("E45:K46")
    # 申込人
    put("I5", "福岡県福岡市中央区赤坂一丁目14番35号\n赤坂門AIビル5階", 9)
    put("I7", "株式会社ふっか")
    put("I9", "代表取締役　松田 圭太", wrap=False)
    # 賃貸物件
    put("F19", f"1階 98.36㎡（約29.75坪）／4階・5階 各229.86㎡（各約69.53坪）\n合計 {T['m2']:.2f}㎡（約{T['tsubo']:.2f}坪）　※1・4・5階セット", 9)
    ws.unmerge_cells("F19:J19"); ws.merge_cells("F19:K19")
    put("F20", "店舗（1階：整形外科クリニック／4・5階：整体院）", 10)
    # 希望賃貸条件
    put("G23", yen(T["rent_total"]), 10, h="right")
    put("I23", "1階 " + yen(T["rent"][0]) + "（97,000円/坪）\n4・5階 各" + yen(T["rent"][1]) + "（45,000円/坪）\n別途消費税等", 8)
    put("G24", yen(T["cam_total"]), 10, h="right")
    put("I24", "1階 " + yen(T["cam"][0]) + "／4・5階 各" + yen(T["cam"][1]) + "\n（坪単価 3,000円/坪）別途消費税等", 8)
    put("G25", yen(T["deposit"]), 10, h="right")
    put("I25", "月額賃料の12ヶ月分\n（1階 " + yen(T["rent"][0] * 12) + "／4・5階 各" + yen(T["rent"][1] * 12) + "）", 8)
    # 特約・要望
    put("D30", "再契約条項追記希望（融資の関係で銀行から依頼がある可能性があるため）\n"
               "株式会社ふっかから一般社団法人福十字へ1階部分を転貸する旨の承諾文言の追記希望"
               "（株式会社ではクリニックの開設・運営ができないため。一般社団法人福十字が診療所の開設者となります）", 9)
    put("D34", "1階・4階・5階のセット賃借希望\nフリーレント3か月希望\n定借10年、6年目以降6か月前予告にて解約可能希望\nMRI設置に伴い、チラー置き場の提供希望", 10)
    # 申込人概要（履歴事項全部証明書より）
    put("E39", "株式会社ふっか（会社法人等番号 2900-01-101980）")
    put("E40", "福岡県福岡市中央区赤坂一丁目14番35号 赤坂門AIビル5階")
    put("J41", "090-5386-3810", h="center")
    put("E42", "代表取締役　松田 圭太")
    put("E43", "2023年（令和5年）6月14日")
    put("E44", "100万円")
    put("E45", "整体治療院の経営（整痛院ふっか）、健康食品・健康器具の企画・販売、運動・食事等による健康事業、"
               "病院・クリニック等の管理・運営に関するコンサルティング、医療技術者の育成・教育 ほか", 9)
    put("E47", "3期分決算書、出店企画書（添付資料・補足資料）、会社概要、会社謄本", 10)
    for r, h in {5: 30, 46: 17.25, 19: 30, 23: 40, 24: 28, 25: 28, 30: 22, 45: 22}.items():
        ws.row_dimensions[r].height = h
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(out)
    print("saved", out, T)
