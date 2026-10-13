"""貸主様式「出店企画書」(xls→xlsx変換済) に記入する。 usage: fill_xlsx.py <template.xlsx> <out.xlsx>"""
import sys
import openpyxl
from openpyxl.styles import Alignment, PatternFill, Font
from geometry import treat_count

src, out = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
ws = wb.active
TODO = PatternFill("solid", fgColor="FFF4C2")
WRAP = Alignment(wrap_text=True, vertical="center")
tc = treat_count()
N = tc["4F"] + tc["5F"]
from finance import model
fin = model()[2]
rent_p, op_p, ad_p = (fin[k] / fin['sales'] * 100 for k in ('rent', 'op', 'ad'))


def put(cell, value, todo=False, merge=None, size=None):
    if merge:
        ws.merge_cells(merge)
    c = ws[cell]
    c.value = value
    c.alignment = WRAP
    if size:
        c.font = Font(name=c.font.name, size=size)
    if todo:
        c.fill = TODO


EST = PatternFill("solid", fgColor="EAF2E3")   # 市場平均による概算（緑）

put("F2", "2026年　10月　12日")
put("E4", "福岡県福岡市中央区赤坂1丁目14番35号 赤坂門AIビル5階", merge="E4:G4", size=9)
put("E5", "株式会社ふっか", merge="E5:G5")
put("E6", "代表取締役　松田 圭太", merge="E6:G6")

put("C8", "1階：整形外科ふっか 東京駅院　／　4・5階：整痛院ふっか 東京駅院", merge="C8:G8", size=10)
put("C9", "1階：整形外科（診察・処置・1.5T MRI検査）\n"
          "4・5階：痛み・しびれ専門の整痛院（手技・運動療法による慢性痛ケア／完全予約制）",
    merge="C9:G9", size=9)
put("C10", "腰痛・肩こり・坐骨神経痛など慢性的な痛み・しびれに悩む40～70代の男女。\n"
           "日本橋・東京駅周辺のオフィスワーカー、近隣居住者、出張・来訪者、銀座院の既存利用者", merge="C10:G10", size=9)
put("C11", "整痛施術（独自のMSMメソッド：手技・運動療法）、MRI検査、診察・処置、セルフケア用品", merge="C11:G11", size=9)

rows = [  # 品目, ％, 特徴, 価格帯（＊は市場平均による概算）
    ("整痛施術（個室）", "80", f"4・5F 個室施術室{N}室。予約制（初回約120分、2回目以降約30分）",
     "痛み消失コース（6ヶ月）880,000円（税込）\n理想の体コース（12ヶ月）1,320,000円（税込）"),
    ("MRI検査", "11", "1F 1.5T超電導MRI（富士フイルム）。痛みの原因を画像で確認",
     "自費診療 30,000～50,000円"),
    ("診察・処置", "7", "1F 診察室・処置室。検査結果の説明・処置",
     "自費診療 3,000～9,000円"),
    ("物販", "2", "セルフケア用品・サポーター等（受付で案内）", "3,000円〜20,000円"),
]
for i, (name, pct, feat, price) in enumerate(rows):
    r = 14 + i
    put(f"B{r}", name, size=9)
    put(f"C{r}", pct + "%", size=9)
    ws[f"C{r}"].fill = EST
    put(f"D{r}", feat, merge=f"D{r}:E{r}", size=8)
    put(f"F{r}", price, merge=f"F{r}:G{r}", size=8)
    if "＊" in price:
        ws[f"F{r}"].fill = EST

put("C19", "銀座院の内装（白い塗り壁・木目・ベージュ系カーペット、中央の円形受付カウンターと木の天井リング）を踏襲し、"
           "規模を拡大して施術室を増設（4F・5F 各10室、各約4畳）。4Fは受付を中央に大きく配置した開放的なロビー。1Fは受付待合・診察室・処置室・MRI検査室を同じ素材感でまとめる。\n"
           "■添付：内観パース（P.9～11）、平面図 A-01（1F・P.3）・A-02（4F・P.4）・A-03（5F・P.5）、仕上表（P.14）",
    size=9)
put("C20", "貸主様のご意向に沿い、看板を目立たせない落ち着いた外観とします。"
           "ビル指定のサイン位置（S1・S2）に店名ロゴのみ表示し、1F正面は木調ルーバー越しの電球色の間接光に留めます。\n"
           "■添付：外観パース（P.8）、立面図 A-04（P.6・サイン位置・使用素材明記）、1F入口詳細図 A-05（P.7）、平面図 A-01",
    size=9)
put("C21", "1F：1.5T超電導MRI（富士フイルム製・ゼロヘリウム型を想定：クエンチ配管不要）、MRI検査室 約31.5㎡（電磁シールド工事）・機械室（専用電源・空調・チラー）、"
          "操作室（検査室との間に観察窓）、MRI更衣室（1名）、診察室・処置室（手洗い器付）の医療機器\n"
           f"4・5F：施術ベッド{N}台、受付カウンター、更衣ロッカー　5F：スタッフ休憩室（ミニキッチン）・スタッフ更衣室・物品庫",
    size=9)
put("C22", "9:00～21:00（最終受付 20:00）予定")
put("C23", "不定休（年末年始を除き原則無休）")
ws["C23"].fill = EST
put("C24", f"最終月商 {fin['sales']:,}万円（3年目・年商 約10億1,700万円）予定　※賃料比率 約{rent_p:.0f}%・営業利益率 約{op_p:.0f}%：別シート「収支計画」／添付資料P.12参照", size=9)
put("C25", "店舗内装イメージ資料、店舗外装イメージ資料\n"
           "（添付資料.pdf：計画概要、平面図A-01～03、立面図・サイン計画A-04、1F入口詳細A-05、外観・内観パース、収支計画、銀座院のご紹介、仕上表・ビル規定対応）", size=9)
put("C26", "松田 圭太\nTEL：090-5386-3810　Mail：kta@medicalbookjapan.co.jp", size=9)
put("C27", "・面積：1F 98.36㎡（約29.75坪）、4F・5F 各G229.86㎡（約69.53坪）、合計約168.8坪（募集資料）。図面寸法は実施設計時にビル側図面で確定します。\n"
           "・本ビルは誘導用途区画のため、診療所・施術所用途の適合性を事前にご確認させていただきます。\n"
           "・MRI（1.5T）は床荷重5,000N/㎡の1Fに配置。重量（荷重分散架台）・搬入経路・シールド工事・電源容量をビル側と協議します。\n"
           f"・賃料比率（最終月商比 約{rent_p:.0f}%）は、駅徒歩1分の立地による広告費の圧縮（売上比 約{ad_p:.0f}%）と自費100%による粗利率 約94%で吸収し、営業利益率 約{op_p:.0f}%を見込みます（収支計画参照）。\n"
           "・緑色セル（＊）は市場平均等による概算値です（構成比・休日）。",
    size=8)

for r, h in {9: 36, 10: 36, 11: 30, 14: 34, 15: 34, 16: 30, 17: 30, 19: 72, 20: 80, 21: 64, 24: 33, 25: 40, 26: 36, 27: 100}.items():
    ws.row_dimensions[r].height = h
ws.page_setup.fitToWidth = 1
ws.page_setup.fitToHeight = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
wb.save(out)
print("saved", out)
